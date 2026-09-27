"""Slide helpers for the Catalyst BSS training deck (python-pptx, 16:9).

Slide kinds: title, section, bullets, two-column, table, worked example
(monospace ledger), screenshot with callouts, quiz. One visual language:
ink on paper, one accent, generous margins, no clip art.
"""
from __future__ import annotations
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INK = RGBColor(0x1A, 0x1D, 0x21)
MUTED = RGBColor(0x5F, 0x66, 0x6E)
ACCENT = RGBColor(0x0B, 0x6E, 0x99)
ACCENT_SOFT = RGBColor(0xE3, 0xF1, 0xF7)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
RULE = RGBColor(0xD9, 0xDD, 0xE1)
OK = RGBColor(0x1E, 0x7F, 0x4F)
WARN = RGBColor(0xB4, 0x5F, 0x06)
FONT = "Calibri"
MONO = "Consolas"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.7)


class Deck:
    def __init__(self, shots_dir: str):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.shots = shots_dir
        self.n = 0
        self.section_name = ""

    # ---- primitives -------------------------------------------------
    def _slide(self):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        return s

    def _text(self, slide, x, y, w, h, text, size=18, bold=False, color=INK, font=FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = font
        return tb

    def _rect(self, slide, x, y, w, h, fill, line=None):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = Pt(0.75)
        shp.shadow.inherit = False
        return shp

    def _chrome(self, slide, title, kicker=None):
        # top rule + kicker + title, bottom footer with section and page
        self._rect(slide, 0, 0, W, Inches(0.12), ACCENT)
        if kicker:
            self._text(slide, M, Inches(0.32), W - 2 * M, Inches(0.35), kicker.upper(), size=11, bold=True, color=ACCENT)
        self._text(slide, M, Inches(0.6), W - 2 * M, Inches(0.8), title, size=28 if len(title) <= 52 else 22, bold=True)
        self._rect(slide, M, Inches(1.38), W - 2 * M, Emu(9525), RULE)
        self._text(slide, M, H - Inches(0.45), W - 2 * M - Inches(1), Inches(0.3), f"Catalyst BSS · training · {self.section_name}", size=10, color=MUTED)
        self._text(slide, W - M - Inches(1), H - Inches(0.45), Inches(1), Inches(0.3), str(self.n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)

    def _bullets(self, slide, x, y, w, h, items, size=17):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        first = True
        for it in items:
            level = 0
            text = it
            if isinstance(it, tuple):
                level, text = it
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.level = level
            p.space_after = Pt(6 if level == 0 else 3)
            bullet = "•" if level == 0 else "–"
            r = p.add_run()
            r.text = f"{bullet}  {text}"
            r.font.size = Pt(size if level == 0 else size - 2)
            r.font.color.rgb = INK if level == 0 else MUTED
            r.font.name = FONT
        return tb

    # ---- slide kinds ------------------------------------------------
    def title(self, title, subtitle, note=None):
        s = self._slide()
        self._rect(s, 0, 0, W, H, PAPER)
        self._rect(s, 0, 0, Inches(0.35), H, ACCENT)
        self._text(s, Inches(1.2), Inches(2.3), W - Inches(2.4), Inches(1.3), title, size=44, bold=True)
        self._text(s, Inches(1.2), Inches(3.6), W - Inches(2.4), Inches(1.0), subtitle, size=22, color=MUTED)
        if note:
            self._text(s, Inches(1.2), H - Inches(1.1), W - Inches(2.4), Inches(0.6), note, size=12, color=MUTED)
        return s

    def section(self, number, name, blurb, agenda=None):
        self.section_name = name
        s = self._slide()
        self._rect(s, 0, 0, W, H, ACCENT)
        self._text(s, Inches(1.2), Inches(1.6), Inches(3), Inches(1), f"{number:02d}", size=64, bold=True, color=PAPER)
        self._text(s, Inches(1.2), Inches(2.75), W - Inches(2.4), Inches(1), name, size=40 if len(name) <= 34 else 30, bold=True, color=PAPER)
        self._text(s, Inches(1.2), Inches(3.85), W - Inches(2.4), Inches(1.1), blurb, size=17, color=ACCENT_SOFT)
        if agenda:
            self._text(s, Inches(1.2), Inches(5.0), W - Inches(2.4), Inches(2.2), ["In this module:  " + "  ·  ".join(agenda)], size=13, color=ACCENT_SOFT)
        return s

    def bullets(self, title, items, kicker=None, note=None):
        s = self._slide()
        self._chrome(s, title, kicker)
        self._bullets(s, M, Inches(1.6), W - 2 * M, Inches(4.9), items)
        if note:
            self._note(s, note)
        return s

    def two_col(self, title, left_title, left, right_title, right, kicker=None, note=None):
        s = self._slide()
        self._chrome(s, title, kicker)
        cw = (W - 2 * M - Inches(0.4)) / 2
        for i, (t, items) in enumerate(((left_title, left), (right_title, right))):
            x = M + i * (cw + Inches(0.4))
            self._rect(s, x, Inches(1.6), cw, Inches(0.45), ACCENT_SOFT)
            self._text(s, x + Inches(0.1), Inches(1.63), cw, Inches(0.4), t, size=14, bold=True, color=ACCENT)
            self._bullets(s, x, Inches(2.15), cw, Inches(4.3), items, size=15)
        if note:
            self._note(s, note)
        return s

    def table(self, title, header, rows, kicker=None, note=None, col_widths=None, size=13):
        s = self._slide()
        self._chrome(s, title, kicker)
        nrows, ncols = len(rows) + 1, len(header)
        top = Inches(1.65)
        height = min(Inches(4.8), Inches(0.42) * nrows)
        gt = s.shapes.add_table(nrows, ncols, M, top, W - 2 * M, height).table
        if col_widths:
            total = sum(col_widths)
            for i, cw in enumerate(col_widths):
                gt.columns[i].width = int((W - 2 * M) * cw / total)
        for c, htxt in enumerate(header):
            cell = gt.cell(0, c)
            cell.text = htxt
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
                    r.font.bold = True
                    r.font.color.rgb = PAPER
                    r.font.name = FONT
        for ri, row in enumerate(rows, start=1):
            for c, val in enumerate(row):
                cell = gt.cell(ri, c)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = PAPER if ri % 2 else RGBColor(0xF4, 0xF6, 0xF8)
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(size)
                        r.font.color.rgb = INK
                        r.font.name = MONO if (isinstance(val, str) and val[:1].isdigit() and any(ch.isdigit() for ch in val)) and c > 0 else FONT
        if note:
            self._note(s, note)
        return s

    def example(self, title, lines, kicker="Worked example", explain=None, note=None):
        """A monospace block (the arithmetic) on the left, the explanation on the right."""
        s = self._slide()
        self._chrome(s, title, kicker)
        bw = Inches(7.7) if explain else W - 2 * M
        longest = max((len(l) for l in lines), default=40)
        msize = 14 if (longest <= 62 and len(lines) <= 11) else (12.5 if longest <= 74 else 11)
        self._rect(s, M, Inches(1.65), bw, Inches(4.75), RGBColor(0xF7, 0xF8, 0xFA), RULE)
        self._text(s, M + Inches(0.2), Inches(1.8), bw - Inches(0.4), Inches(4.5), lines, size=msize, font=MONO)
        if explain:
            x = M + bw + Inches(0.35)
            self._bullets(s, x, Inches(1.65), W - M - x, Inches(4.75), explain, size=15)
        if note:
            self._note(s, note)
        return s

    def shot(self, title, image, callouts=None, kicker="Where to click", note=None):
        """Screenshot (scaled into the frame) with numbered callouts listed at the right."""
        s = self._slide()
        self._chrome(s, title, kicker)
        path = os.path.join(self.shots, image)
        iw_max = Inches(8.9) if callouts else W - 2 * M
        ih_max = Inches(4.9)
        if os.path.exists(path):
            from PIL import Image
            with Image.open(path) as im:
                w, h = im.size
            scale = min(iw_max / w, ih_max / h)
            iw, ih = int(w * scale), int(h * scale)
            pic = s.shapes.add_picture(path, M, Inches(1.6), iw, ih)
            pic.line.color.rgb = RULE
            pic.line.width = Pt(0.75)
        else:
            self._rect(s, M, Inches(1.6), iw_max, ih_max, RGBColor(0xF4, 0xF6, 0xF8), RULE)
            self._text(s, M, Inches(3.8), iw_max, Inches(0.5), f"screenshot pending: {image}", size=14, color=MUTED, align=PP_ALIGN.CENTER)
        if callouts:
            x = M + iw_max + Inches(0.3)
            items = [f"{i+1}. {c}" for i, c in enumerate(callouts)]
            tb = s.shapes.add_textbox(x, Inches(1.6), W - M - x, Inches(4.9))
            tf = tb.text_frame
            tf.word_wrap = True
            for i, it in enumerate(items):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(8)
                r = p.add_run()
                r.text = it
                r.font.size = Pt(13)
                r.font.color.rgb = INK
                r.font.name = FONT
        if note:
            self._note(s, note)
        return s

    def quiz(self, title, questions, answers=None):
        s = self._slide()
        self._chrome(s, title, "Check yourself")
        items = [f"Q{i+1}. {q}" for i, q in enumerate(questions)]
        self._bullets(s, M, Inches(1.6), W - 2 * M if not answers else Inches(6.4), Inches(4.9), items, size=15)
        if answers:
            x = M + Inches(6.7)
            self._rect(s, x, Inches(1.6), W - M - x, Inches(4.9), ACCENT_SOFT)
            self._text(s, x + Inches(0.15), Inches(1.7), W - M - x - Inches(0.3), Inches(0.4), "Answers", size=13, bold=True, color=ACCENT)
            self._bullets(s, x + Inches(0.1), Inches(2.1), W - M - x - Inches(0.2), Inches(4.3), [f"A{i+1}. {a}" for i, a in enumerate(answers)], size=12)
        return s

    def _note(self, slide, note):
        self._rect(slide, M, Inches(6.45), W - 2 * M, Inches(0.5), ACCENT_SOFT)
        self._text(slide, M + Inches(0.15), Inches(6.5), W - 2 * M - Inches(0.3), Inches(0.42), note, size=12, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)

    def save(self, path):
        self.prs.save(path)
        return self.n
